
from io import BytesIO
import secrets
import string

from django import forms
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render
from django.urls import reverse_lazy
from django.views.generic import CreateView, FormView, ListView, UpdateView, View

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import PerfilTecnico
from .mixins import RolRequeridoMixin
from .roles import ADMINISTRADOR
from .forms import CargaMasivaTecnicosForm


User = get_user_model()
GRUPO_TECNICO = "Tecnico"


def error_403(request, exception=None):
    return render(request, "403.html", status=403)


def obtener_grupo_tecnico():
    grupo, _ = Group.objects.get_or_create(name=GRUPO_TECNICO)
    return grupo


# =========================================================
# FORMULARIO DE CREACIÓN Y EDICIÓN
# =========================================================

class TecnicoForm(forms.Form):
    username = forms.CharField(
        label="Usuario",
        max_length=150,
    )
    first_name = forms.CharField(
        label="Apellidos y nombres",
        max_length=150,
    )
    email = forms.EmailField(
        label="Correo electrónico",
        required=False,
    )
    cedula = forms.CharField(
        label="Cédula",
        max_length=20,
    )
    supervisor = forms.CharField(
        label="Supervisor",
        max_length=150,
        required=False,
    )

    def __init__(self, *args, user_instance=None, **kwargs):
        self.user_instance = user_instance
        super().__init__(*args, **kwargs)

    def clean_username(self):
        username = self.cleaned_data["username"].strip()

        usuarios = User.objects.filter(username__iexact=username)

        if self.user_instance:
            usuarios = usuarios.exclude(pk=self.user_instance.pk)

        if usuarios.exists():
            raise forms.ValidationError(
                "Este nombre de usuario ya está registrado."
            )

        return username

    def clean_cedula(self):
        cedula = self.cleaned_data["cedula"].strip()

        if not cedula.isdigit():
            raise forms.ValidationError(
                "La cédula debe contener únicamente números."
            )

        perfiles = PerfilTecnico.objects.filter(cedula=cedula)

        if self.user_instance:
            perfiles = perfiles.exclude(
                usuario=self.user_instance
            )

        if perfiles.exists():
            raise forms.ValidationError(
                "Esta cédula ya está registrada."
            )

        return cedula


# =========================================================
# LISTADO DE TÉCNICOS
# =========================================================

class TecnicoListView(RolRequeridoMixin, ListView):
    model = User
    template_name = "usuarios/tecnico_list.html"
    context_object_name = "tecnicos"
    paginate_by = 20

    def get_queryset(self):
        qs = User.objects.filter(
            groups__name=GRUPO_TECNICO,
            perfil_tecnico__isnull=False,
        ).select_related(
            "perfil_tecnico"
        ).distinct()

        q = self.request.GET.get("q", "").strip()

        if q:
            qs = qs.filter(
                Q(username__icontains=q)
                | Q(first_name__icontains=q)
                | Q(last_name__icontains=q)
                | Q(email__icontains=q)
                | Q(perfil_tecnico__cedula__icontains=q)
                | Q(perfil_tecnico__supervisor__icontains=q)
            )

        return qs.order_by("first_name", "last_name", "username")

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)

        params = self.request.GET.copy()
        params.pop("page", None)

        ctx["qs"] = params.urlencode()

        return ctx


# =========================================================
# CREAR Y EDITAR TÉCNICOS
# =========================================================

class TecnicoFormMixin(RolRequeridoMixin):
    roles_permitidos = (ADMINISTRADOR,)
    form_class = TecnicoForm
    template_name = "form.html"
    success_url = reverse_lazy("tecnicos:lista")

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()

        if isinstance(self, UpdateView):
            kwargs["user_instance"] = self.get_object()

        return kwargs

    def get_initial(self):
        initial = super().get_initial()

        if isinstance(self, UpdateView):
            usuario = self.get_object()
            initial.update({
                "username": usuario.username,
                "first_name": usuario.first_name,
                "email": usuario.email,
                "cedula": usuario.perfil_tecnico.cedula,
                "supervisor": usuario.perfil_tecnico.supervisor,
            })

        return initial

    def form_valid(self, form):
        datos = form.cleaned_data

        with transaction.atomic():
            if isinstance(self, UpdateView):
                usuario = self.get_object()
                perfil = usuario.perfil_tecnico
            else:
                usuario = User()

                # No asignar contraseñas conocidas o predecibles.
                usuario.set_unusable_password()
                perfil = PerfilTecnico(usuario=usuario)

            usuario.username = datos["username"]
            usuario.first_name = datos["first_name"]
            usuario.email = datos["email"]
            usuario.save()

            perfil.usuario = usuario
            perfil.cedula = datos["cedula"]
            perfil.supervisor = datos["supervisor"]
            perfil.save()

            usuario.groups.add(obtener_grupo_tecnico())

        messages.success(
            self.request,
            "Técnico guardado correctamente.",
        )

        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["cancel_url"] = self.success_url
        return ctx


class TecnicoCreateView(TecnicoFormMixin, CreateView):
    extra_context = {
        "titulo": "Nuevo técnico",
        "boton": "Guardar técnico",
    }


class TecnicoUpdateView(TecnicoFormMixin, UpdateView):
    model = User
    queryset = User.objects.filter(
        groups__name=GRUPO_TECNICO,
        perfil_tecnico__isnull=False,
    ).select_related("perfil_tecnico")
    extra_context = {
        "titulo": "Editar técnico",
        "boton": "Guardar cambios",
    }


# =========================================================
# DESCARGAR PLANTILLA EXCEL
# =========================================================

class DescargarPlantillaTecnicosView(RolRequeridoMixin, View):
    roles_permitidos = (ADMINISTRADOR,)

    def get(self, request, *args, **kwargs):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Técnicos"

        encabezados = [
            "SUPERVISOR",
            "CEDULA",
            "APELLIDOS Y NOMBRES",
        ]
        worksheet.append(encabezados)

        for columna, encabezado in enumerate(encabezados, start=1):
            celda = worksheet.cell(row=1, column=columna)
            celda.fill = PatternFill(
                fill_type="solid",
                fgColor="2D6A4F",
            )
            celda.font = Font(color="FFFFFF", bold=True)
            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        worksheet.column_dimensions["A"].width = 30
        worksheet.column_dimensions["B"].width = 20
        worksheet.column_dimensions["C"].width = 40
        worksheet.row_dimensions[1].height = 28
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = "A1:C1000"

        for fila in range(2, 1001):
            worksheet.cell(row=fila, column=2).number_format = "@"

        instrucciones = workbook.create_sheet("Instrucciones")
        instrucciones["A1"] = "CARGA MASIVA DE TÉCNICOS"
        instrucciones["A1"].font = Font(
            bold=True,
            size=16,
            color="2D6A4F",
        )
        instrucciones["A3"] = "COLUMNAS"
        instrucciones["A3"].font = Font(bold=True)
        instrucciones["A4"] = "SUPERVISOR"
        instrucciones["B4"] = "Nombre del supervisor del técnico."
        instrucciones["A5"] = "CEDULA"
        instrucciones["B5"] = "Número de cédula del técnico."
        instrucciones["A6"] = "APELLIDOS Y NOMBRES"
        instrucciones["B6"] = "Nombre completo del técnico."
        instrucciones["A8"] = "IMPORTANTE"
        instrucciones["A8"].font = Font(bold=True, color="C62828")
        instrucciones["A9"] = "No cambie los nombres de las columnas."
        instrucciones["A10"] = "La cédula debe ser única."
        instrucciones["A11"] = (
            "Se crearán usuarios con contraseña no utilizable "
            "hasta que se establezca una nueva."
        )
        instrucciones.column_dimensions["A"].width = 30
        instrucciones.column_dimensions["B"].width = 75

        archivo = BytesIO()
        workbook.save(archivo)
        archivo.seek(0)

        response = HttpResponse(
            archivo.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )
        response["Content-Disposition"] = (
            'attachment; filename="plantilla_tecnicos.xlsx"'
        )
        return response


# =========================================================
# CARGA MASIVA DE TÉCNICOS
# =========================================================

class CargaMasivaTecnicosView(RolRequeridoMixin, FormView):
    roles_permitidos = (ADMINISTRADOR,)
    template_name = "usuarios/tecnico_carga_masiva.html"
    form_class = CargaMasivaTecnicosForm
    success_url = reverse_lazy("tecnicos:lista")

    def form_valid(self, form):
        archivo = form.cleaned_data["archivo"]

        try:
            workbook = load_workbook(
                archivo,
                read_only=True,
                data_only=True,
            )
            worksheet = workbook.active
        except Exception:
            form.add_error(
                "archivo",
                "No fue posible leer el archivo Excel.",
            )
            return self.form_invalid(form)

        try:
            encabezados_esperados = [
                "SUPERVISOR",
                "CEDULA",
                "APELLIDOS Y NOMBRES",
            ]

            encabezados_excel = [
                str(celda.value).strip().upper()
                if celda.value is not None else ""
                for celda in worksheet[1]
            ]

            while encabezados_excel and encabezados_excel[-1] == "":
                encabezados_excel.pop()

            if encabezados_excel != encabezados_esperados:
                form.add_error(
                    "archivo",
                    "Las columnas deben ser exactamente: "
                    "SUPERVISOR, CEDULA y APELLIDOS Y NOMBRES.",
                )
                return self.form_invalid(form)

            registros = []
            cedulas_excel = set()
            errores = []

            for numero_fila, fila in enumerate(
                worksheet.iter_rows(min_row=2, values_only=True),
                start=2,
            ):
                fila = tuple(fila[:3])

                if not any(
                    valor is not None and str(valor).strip()
                    for valor in fila
                ):
                    continue

                supervisor = (
                    str(fila[0]).strip()
                    if fila[0] is not None else ""
                )
                cedula = (
                    str(fila[1]).strip()
                    if fila[1] is not None else ""
                )
                nombre = (
                    str(fila[2]).strip()
                    if fila[2] is not None else ""
                )

                if not supervisor:
                    errores.append(
                        f"Fila {numero_fila}: falta el supervisor."
                    )
                    continue

                if not cedula or not cedula.isdigit():
                    errores.append(
                        f"Fila {numero_fila}: la cédula debe "
                        "contener únicamente números."
                    )
                    continue

                if len(cedula) > 20:
                    errores.append(
                        f"Fila {numero_fila}: la cédula supera "
                        "los 20 caracteres."
                    )
                    continue

                if not nombre:
                    errores.append(
                        f"Fila {numero_fila}: falta el nombre."
                    )
                    continue

                if cedula in cedulas_excel:
                    errores.append(
                        f"Fila {numero_fila}: cédula duplicada "
                        f"en el Excel ({cedula})."
                    )
                    continue

                cedulas_excel.add(cedula)
                registros.append({
                    "supervisor": supervisor,
                    "cedula": cedula,
                    "nombre": nombre,
                })

            if errores:
                for error in errores:
                    form.add_error("archivo", error)
                return self.form_invalid(form)

            if not registros:
                form.add_error(
                    "archivo",
                    "El Excel no contiene técnicos para importar.",
                )
                return self.form_invalid(form)

            cedulas_existentes = set(
                PerfilTecnico.objects.filter(
                    cedula__in=cedulas_excel
                ).values_list("cedula", flat=True)
            )

            if cedulas_existentes:
                form.add_error(
                    "archivo",
                    "Estas cédulas ya están registradas: "
                    + ", ".join(sorted(cedulas_existentes)),
                )
                return self.form_invalid(form)

            # Evitar sobrescribir cuentas existentes.
            usernames = [registro["cedula"] for registro in registros]
            usuarios_existentes = set(
                User.objects.filter(
                    username__in=usernames
                ).values_list("username", flat=True)
            )

            if usuarios_existentes:
                form.add_error(
                    "archivo",
                    "Ya existen usuarios con estos nombres de usuario: "
                    + ", ".join(sorted(usuarios_existentes)),
                )
                return self.form_invalid(form)

            with transaction.atomic():
                grupo = obtener_grupo_tecnico()

                for registro in registros:
                    usuario = User(
                        username=registro["cedula"],
                        first_name=registro["nombre"],
                        is_active=True,
                    )
                    usuario.set_unusable_password()
                    usuario.save()
                    usuario.groups.add(grupo)

                    PerfilTecnico.objects.create(
                        usuario=usuario,
                        cedula=registro["cedula"],
                        supervisor=registro["supervisor"],
                    )

        except Exception:
            # No ocultar los detalles de validación como errores de archivo.
            # La excepción inesperada se propaga para poder diagnosticarla.
            raise
        finally:
            workbook.close()

        messages.success(
            self.request,
            f"Se cargaron correctamente {len(registros)} técnico(s). "
            "Los usuarios deben establecer sus contraseñas antes de "
            "iniciar sesión.",
        )

        return super().form_valid(form)
